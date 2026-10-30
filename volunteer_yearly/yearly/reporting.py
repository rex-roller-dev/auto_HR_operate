from decimal import Decimal
from html import escape
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import hours_text, year_label


def text_cell(cell, value):
    cell.value = str(value)
    cell.data_type = "s"  # 明确文本，避免活动名称等被当作公式执行。


def write_xlsx(report, destination):
    wb = Workbook()
    summary = wb.active
    summary.title = "学年汇总"
    summary.append(["志愿服务学年统计", report.submission.name, report.submission.student_id])
    text_cell(summary.cell(1, 2), report.submission.name)
    text_cell(summary.cell(1, 3), report.submission.student_id)
    summary.append(["评选学年", year_label(report.evaluation_year), "统计截止日", report.as_of.isoformat()])
    summary.append(["统计口径", "按活动结束日期归属学年；核定时长全额计入；i志愿始终排除志愿深圳来源；不计培训时长"])
    summary.append(["学年", "起始日期", "结束日期", "志愿深圳（小时）", "i志愿（小时）", "合计（小时）", "说明"])
    years = sorted(set(report.totals) | set(range(report.evaluation_year - 3, report.evaluation_year + 1)))
    for year in years:
        if year < report.submission.enrollment:
            summary.append([year_label(year), f"{year}-09-01", f"{year+1}-08-31", "未入学", "未入学", "未入学", "早于入学学年"])
            continue
        total = report.totals.get(year)
        finish = min(report.as_of.isoformat(), f"{year+1}-08-31")
        summary.append([
            year_label(year), f"{year}-09-01", finish,
            float(total.shenzhen / 60) if total else 0,
            float(total.ivolunteer / 60) if total else 0,
            float(total.total / 60) if total else 0,
            "评选学年" if year == report.evaluation_year else "",
        ])
    grand = sum((y.total for y in report.totals.values()), Decimal(0))
    summary.append(["入学至截止日合计", "", "", "", "", float(grand / 60)])
    for row in summary.iter_rows(min_row=5):
        for cell in row[3:6]:
            if cell.data_type == "n":
                cell.number_format = "0.00"
    detail = wb.create_sheet("活动明细")
    detail.append(["序号", "平台", "数据来源", "活动名称", "开始日期时间", "结束日期时间", "所属学年",
                   "核定时长（小时）", "计入时长（小时）", "是否计入", "排除原因", "考勤编号", "原附件", "页码", "时长类型"])
    for index, item in enumerate(report.details, 1):
        a = item.activity
        detail.append([index, a.platform, a.source, a.title, a.begin.isoformat(" "), a.end.isoformat(" "),
                       year_label(a.year), float(a.minutes / 60), float(a.minutes / 60) if item.counted else 0,
                       "是" if item.counted else "否", item.reason, a.record_id, a.filename, a.page, a.kind])
        for col in (2, 3, 4, 5, 6, 7, 10, 11, 12, 13, 15):
            text_cell(detail.cell(index + 1, col), detail.cell(index + 1, col).value or "")
        for col in (8, 9):
            detail.cell(index + 1, col).number_format = "0.00####"
    notes = wb.create_sheet("统计说明")
    for row in [
        ["事项", "说明"], ["材料范围", "统计仅覆盖实际上传的PDF记录；未上传的平台显示0，不表示其真实总时长为0。"],
        ["学年", "每年9月1日至次年8月31日，活动全部核定时长按结束日期归属。"],
        ["精度", "内部用Decimal分钟累加，汇总显示两位小数；明细保留更多小数以便复核。"],
        ["去重", "志愿深圳按考勤编号识别重复；i志愿按活动名称、日期、来源与时长识别完全重复记录。"],
        ["来源排除", "无论是否上传志愿深圳PDF，i志愿中来源含“志愿深圳”的记录均不计入。"],
        ["未入学", "早于学号前四位所对应入学学年的记录不计入，评选汇总单元格标记“未入学”。"],
        ["明细", "被排除的记录仍保留，注明原因；原PDF中的联系人、手机号、证件号码不进入统计文件。"],
        *[["提示", note] for note in report.warnings],
    ]:
        notes.append(row)
    for sheet, header_row in ((summary, 4), (detail, 1), (notes, 1)):
        sheet.freeze_panes = f"A{header_row+1}"
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        for c in sheet[header_row]:
            c.font = Font(color="FFFFFF", bold=True)
            c.fill = PatternFill("solid", fgColor="1D4E78")
        for row in sheet.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 22
    detail.column_dimensions["D"].width = 55
    detail.column_dimensions["K"].width = 38
    detail.column_dimensions["M"].width = 35
    notes.column_dimensions["B"].width = 100
    summary.column_dimensions["B"].width = 24
    summary.row_dimensions[3].height = 42
    wb.save(destination)
    return Path(destination)


def email_bodies(report):
    rows = []
    plain = [f"{report.submission.name} 同学，您好：", "", f"评选学年：{year_label(report.evaluation_year)}",
             f"统计截止日期：{report.as_of.isoformat()}", ""]
    for year, value in zip(range(report.evaluation_year, report.evaluation_year - 4, -1), report.evaluation_cells()):
        rows.append(f"<tr><td>{year_label(year)}</td><td>{escape(value)}</td></tr>")
        plain.append(f"{year_label(year)}：{value}" + ("" if value == "未入学" else " 小时"))
    note = "本表仅展示评选学年及前三个学年的志愿服务时长，不生成个人活动明细附件。按活动结束日期归属学年；i志愿中来自志愿深圳的记录始终排除。统计仅基于上传材料，未上传的平台不代表真实时长为0。"
    plain.extend(["", note])
    plain.extend(report.warnings)
    html = (
        f"<p>{escape(report.submission.name)} 同学，您好：</p>"
        f"<p>评选学年：{year_label(report.evaluation_year)}；统计截止：{report.as_of.isoformat()}</p>"
        '<table border="1" cellspacing="0" cellpadding="8"><tr><th>学年</th><th>服务时长（小时）</th></tr>'
        + "".join(rows) + f"</table><p>{note}</p>"
        + "".join(f"<p>{escape(w)}</p>" for w in report.warnings)
    )
    return "\n".join(plain), html
