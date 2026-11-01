from datetime import date
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Table, TableStyle, Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors

from yearly.models import SZ, IV, aggregate
from yearly.parsers import parse_pdf
from yearly.reporting import write_xlsx, email_bodies
from tests.test_rules import student, activity

pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))


def make_pdf(path, platform=SZ, watermark=False, no_grid=False, source="i志愿"):
    width, height = landscape(A4)
    c = canvas.Canvas(str(path), pagesize=(width, height))
    c.setFont("STSong-Light", 11)
    c.drawString(30, height - 30, "姓名 测试同学" if platform == SZ else "志愿者:测试同学")
    if no_grid:
        center = height - 140
        for x, y, text in [
            (40, center, "测试跨学年活动"), (220, center + 17, "2025.08.31"),
            (250, center, "至"), (220, center - 17, "2025.09.01"),
            (665, center, "8时11分"), (750, center, source),
        ]:
            c.drawString(x, y, text)
    else:
        if platform == SZ:
            data = [
                ["考勤编号", "活动名称", "服务时间", "服务时长", "证明人", "时长类型"],
                ["123456", "测试跨学年活动", "2025-08-31 23:00<br/>2025-09-01 01:00", "1.25", "测试", "服务时长"],
                ["123457", "培训活动", "2025-10-01 10:00<br/>2025-10-01 12:00", "2", "测试", "培训时长"],
            ]
            widths = [80, 220, 140, 70, 70, 70]
        else:
            data = [
                ["志愿组织", "活动名称", "活动时间", "志愿团体", "负责人", "联系人", "联系电话", "服务时间", "数据来源"],
                ["", "跨学年测试活动", "2025.08.31<br/>至<br/>2025.09.01", "测试团体", "测试", "测试", "000", "8时11分", "i志愿"],
                ["", "需排除的深圳活动", "2025.09.02<br/>至<br/>2025.09.03", "测试团体", "测试", "测试", "000", "2时30分", "志愿深圳"],
            ]
            widths = [30, 150, 110, 75, 40, 40, 80, 65, 60]
        style = ParagraphStyle("CJK", fontName="STSong-Light", fontSize=10, leading=14)
        data = [[Paragraph(str(cell), style) for cell in row] for row in data]
        table = Table(data, colWidths=widths)
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                                   ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        _, h = table.wrap(780, height)
        table.drawOn(c, 30, height - 70 - h)
    if watermark:
        c.saveState()
        c.translate(260, 100)
        c.rotate(35)
        c.drawString(0, 0, "www.sva.org.cn")
        c.restoreState()
    c.save()


def make_date_pdf(path, platform=SZ, date_values=("-",), durations=None, no_grid=False):
    durations = durations or (["1"] * len(date_values) if platform == SZ else ["1时"] * len(date_values))
    width, height = landscape(A4)
    c = canvas.Canvas(str(path), pagesize=(width, height))
    c.setFont("STSong-Light", 11)
    c.drawString(30, height - 30, "姓名 测试同学" if platform == SZ else "志愿者:测试同学")
    if no_grid:
        c.drawString(220, height - 65, "活动时间")
        for index, (value, duration) in enumerate(zip(date_values, durations)):
            center = height - 140 - index * 90
            c.drawString(40, center, "志愿活动")
            lines = value.splitlines() or [""]
            for line_index, line in enumerate(lines):
                c.drawString(220 if line != "至" else 250, center + (len(lines)-1) * 8.5 - line_index * 17, line)
            c.drawString(665, center, duration)
            c.drawString(750, center, IV)
    else:
        if platform == SZ:
            data = [["考勤编号", "活动名称", "活动时间", "服务时长", "证明人", "时长类型"]]
            data.extend([[str(1000 + i), "2021-2022学年志愿活动", value.replace("\n", "<br/>"),
                          duration, "老师", "服务时长"]
                         for i, (value, duration) in enumerate(zip(date_values, durations))])
            widths = [70, 220, 140, 70, 60, 70]
        else:
            data = [["志愿组织", "活动名称", "活动时间", "志愿地点", "负责人", "联系人", "联系电话", "服务时间", "数据来源"]]
            data.extend([["", "志愿活动", value.replace("\n", "<br/>"), "校园", "老师", "老师", "000", duration, IV]
                         for value, duration in zip(date_values, durations)])
            widths = [30, 150, 110, 75, 40, 40, 80, 65, 60]
        style = ParagraphStyle("DateRows", fontName="STSong-Light", fontSize=10, leading=14)
        table = Table([[Paragraph(str(cell), style) for cell in row] for row in data], colWidths=widths)
        table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                                   ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        _, table_height = table.wrap(780, height)
        table.drawOn(c, 30, height - 70 - table_height)
    c.save()


class PDFTests(unittest.TestCase):
    def test_shenzhen_dash_blank_and_unknown_dates_still_retain_service_hours(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sz.pdf"
            make_date_pdf(path, date_values=("-", "", "—", "暂无"), durations=("12", "12", "12", "4.5"))
            parsed = parse_pdf(path, SZ)
        self.assertEqual(len(parsed.activities), 4)
        self.assertTrue(all(not row.has_complete_dates and row.year is None for row in parsed.activities))
        report = aggregate(student(), parsed.activities, 2025, date(2026, 8, 31))
        self.assertEqual(report.material_hours, "40.50")
        self.assertEqual(report.undated_hours, "40.50")
        self.assertEqual(report.undated_count, 4)
        self.assertEqual(report.evaluation_cells(), ["0.00", "0.00"])
        plain, html = email_bodies(report)
        self.assertIn("40.50 小时志愿服务记录（4条）", plain)
        self.assertIn("未计入任何学年", html)

    def test_mixed_shenzhen_dates_preserve_cutoff_and_undated_totals(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sz.pdf"
            make_date_pdf(path, date_values=("2026-08-30 08:00\n2026-08-31 18:00", "-",
                                            "2026-08-31 08:00\n2026-09-01 18:00"), durations=("2", "3", "99"))
            parsed = parse_pdf(path, SZ)
        report = aggregate(student(), parsed.activities, 2025, date(2026, 8, 31))
        self.assertEqual(report.material_hours, "5.00")
        self.assertEqual(report.evaluation_cells()[0], "2.00")
        self.assertEqual(report.undated_hours, "3.00")

    def test_invalid_dates_and_durations_are_not_accepted_as_missing(self):
        for platform, date_value, duration in ((SZ, "2026-02-30\n-", "2"), (SZ, "错误日期", "2"),
                                                (SZ, "-", "-1"), (SZ, "-", "NaN"), (IV, "-", "abc")):
            with self.subTest(platform=platform, date_value=date_value, duration=duration), tempfile.TemporaryDirectory() as folder:
                path = Path(folder) / "invalid.pdf"
                make_date_pdf(path, platform, (date_value,), (duration,))
                with self.assertRaises(ValueError):
                    parse_pdf(path, platform)

    def test_ivolunteer_table_retains_undated_rows_alongside_dated_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "iv.pdf"
            make_date_pdf(path, IV, ("2025.08.31\n至\n2025.09.01", "-", ""), ("1时", "2时", "3时"))
            parsed = parse_pdf(path, IV)
        self.assertEqual(len(parsed.activities), 3)
        report = aggregate(student(), parsed.activities, 2025, date(2026, 8, 31))
        self.assertEqual(report.material_hours, "6.00")
        self.assertEqual(report.evaluation_cells()[0], "1.00")
        self.assertEqual(report.undated_hours, "5.00")

    def test_identical_undated_ivolunteer_rows_remain_distinct_and_repeat_pdf_is_deduplicated(self):
        for no_grid in (False, True):
            with self.subTest(no_grid=no_grid), tempfile.TemporaryDirectory() as folder:
                path = Path(folder) / "iv.pdf"
                make_date_pdf(path, IV, ("-", ""), ("2时", "2时"), no_grid=no_grid)
                parsed = parse_pdf(path, IV)
                self.assertEqual(len(parsed.activities), 2)
                self.assertNotEqual(parsed.activities[0].identity, parsed.activities[1].identity)
                report = aggregate(student(), parsed.activities * 2, 2025, date(2026, 8, 31))
                self.assertEqual(report.material_hours, "4.00")
                self.assertEqual(report.undated_count, 2)

    def test_borderless_ivolunteer_mixed_dates_and_missing_date(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "iv.pdf"
            make_date_pdf(path, IV, ("2025.08.31\n至\n2025.09.01", "-",
                                    "2026.08.31\n至\n2026.09.01"), ("1时", "2时", "99时"), no_grid=True)
            parsed = parse_pdf(path, IV)
        self.assertEqual(len(parsed.activities), 3)
        report = aggregate(student(), parsed.activities, 2025, date(2026, 8, 31))
        self.assertEqual(report.material_hours, "3.00")
        self.assertEqual(report.undated_hours, "2.00")
        self.assertEqual(report.evaluation_cells()[0], "1.00")

    def test_shenzhen_real_pdf_tables_and_watermark(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sz.pdf"
            make_pdf(path, SZ, watermark=True)
            parsed = parse_pdf(path, SZ)
        self.assertEqual(len(parsed.activities), 2)
        self.assertEqual(parsed.names, {"测试同学"})
        self.assertEqual(parsed.activities[0].minutes, Decimal(75))
        self.assertEqual(parsed.activities[1].kind, "培训时长")

    def test_ivolunteer_real_pdf_table_retains_end_dates_and_sources(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "iv.pdf"
            make_pdf(path, IV)
            parsed = parse_pdf(path, IV)
        self.assertEqual(len(parsed.activities), 2)
        self.assertEqual(parsed.activities[0].end.date(), date(2025, 9, 1))
        self.assertEqual(parsed.activities[0].minutes, Decimal(491))
        report = aggregate(student(), parsed.activities, 2025, date(2026, 10, 30))
        self.assertEqual(report.evaluation_cells()[0], "8.18")
        self.assertIn("来源为志愿深圳", report.details[1].reason)

    def test_ivolunteer_without_borders_uses_separate_record_band(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "iv.pdf"
            make_pdf(path, IV, no_grid=True)
            parsed = parse_pdf(path, IV)
        self.assertEqual(len(parsed.activities), 1)
        self.assertEqual(parsed.activities[0].minutes, Decimal(491))
        self.assertEqual(parsed.activities[0].end.date(), date(2025, 9, 1))

    def test_empty_pdf_not_silently_zero(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "blank.pdf"
            c = canvas.Canvas(str(path))
            c.showPage()
            c.save()
            with self.assertRaises(ValueError):
                parse_pdf(path, IV)

    def test_ivolunteer_accepts_guangzhou_volunteer_source(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "iv.pdf"
            make_pdf(path, IV, no_grid=True, source="广州志愿时")
            parsed = parse_pdf(path, IV)
        self.assertEqual(len(parsed.activities), 1)
        self.assertEqual(parsed.activities[0].source, "广州志愿时")
        self.assertEqual(parsed.activities[0].minutes, Decimal(491))

    def test_email_contains_all_enrolled_years_and_material_total_column(self):
        report = aggregate(student("2020"), [activity(rid="now"),
                           activity("2019-10-01", "2019-10-02", "3", rid="before")],
                           2025, date(2026, 8, 31))
        plain, html = email_bodies(report)
        for year in ("2025-2026", "2024-2025", "2023-2024", "2022-2023", "2021-2022", "2020-2021"):
            self.assertEqual(html.count("<td>" + year + "</td>"), 1)
        self.assertNotIn("<td>2019-2020</td>", html)
        self.assertNotIn("2026-2027", html)
        self.assertIn("<th>材料总时长（小时）</th>", html)
        self.assertIn('<td rowspan="6">7.00</td>', html)
        self.assertIn("材料总时长：7.00 小时", plain)
        self.assertNotIn("附件包含", plain)
        self.assertNotIn("活动名称", html)

    def test_email_total_still_shown_if_enrollment_is_after_evaluation_year(self):
        report = aggregate(student("2026"), [activity()], 2025, date(2026, 8, 31))
        plain, html = email_bodies(report)
        self.assertIn("截至评选学年尚未入学", plain)
        self.assertIn("材料总时长：4.00 小时", plain)
        self.assertNotIn('rowspan="0"', html)
        self.assertNotIn("<td>2026-2027</td>", html)

    def test_workbook_contains_summary_details_and_explicit_text(self):
        a = activity()
        from dataclasses import replace
        a = replace(a, title='=HYPERLINK("https://example.com")')
        report = aggregate(student(), [a], 2025, date(2026, 10, 30))
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "report.xlsx"
            write_xlsx(report, path)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["学年汇总", "活动明细", "统计说明"])
            self.assertEqual(wb["活动明细"]["D2"].data_type, "s")
            self.assertEqual(wb["活动明细"]["I2"].value, 4)
            self.assertIn("未入学", [c.value for row in wb["学年汇总"] for c in row])
